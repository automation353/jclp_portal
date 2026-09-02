"""REST endpoints for the PO Data module.

  POST /api/po-data/uploads/new/  accept a PO file, save + forward to n8n
  GET  /api/po-data/uploads/      list the last 30 uploads

Parses the PO xlsx on the Django side (openpyxl, already installed) and
sends TWO row-lists to n8n:

  * ``po``                     — every row with every column
  * ``po_highlighted_columns`` — the same rows restricted to the columns
                                  listed in PO_HIGHLIGHTED_COLUMNS below

If PO_HIGHLIGHTED_COLUMNS is None, the subset payload matches ``po`` 1:1.
Edit the list here when you want the ``po_highlighted_columns`` tab
in the target Google Sheet to carry a specific subset.

No Google Apps Script needed — the n8n workflow writes both tabs directly.
"""

import logging
import os

import requests
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.decorators import api_view, parser_classes
from rest_framework.parsers import MultiPartParser
from rest_framework.response import Response

from .models import POUpload
from .serializers import POUploadSerializer


log = logging.getLogger(__name__)

UPLOAD_ROOT_DEFAULT = "/root/jclp_automation_portal/jcpl/uploads/po_data"

TARGET_SHEET_URL = (
    "https://docs.google.com/spreadsheets/d/"
    "1jg4EFvWaRjwJ4l8tt0FqBM0Fy3sHT-0y6LrWhYizMQ0/edit"
)

# If a sheet named exactly one of these exists in the workbook, we read from
# it. Otherwise we fall back to the workbook's first sheet.
PREFERRED_SHEET_NAMES = ("PO", "Purchase Order", "po", "purchase_order")

# The header row is auto-detected (first row with ≥ 3 non-empty cells). If
# your PO file has a specific header row you always want, set this to a
# positive integer (1-indexed). Leave as None for auto-detect.
PO_HEADER_ROW_OVERRIDE = None

# List the exact column-header strings you want mirrored to the
# ``po_highlighted_columns`` tab. None → mirror every column (same as
# the ``po`` tab). Empty list → the tab will still exist but with no rows.
#
# When you tell me the exact PO column names you want highlighted, drop
# them in this list — no rebuild needed, next upload uses them.
PO_HIGHLIGHTED_COLUMNS = None


def _upload_root():
    return os.environ.get("JCLP_PO_UPLOAD_ROOT", UPLOAD_ROOT_DEFAULT)


def _webhook_url():
    return os.environ.get("JCLP_PO_SHEET_WEBHOOK", "").strip()


def _store_uploaded_file(uploaded_file) -> tuple:
    now = timezone.now()
    subdir = os.path.join(_upload_root(), now.strftime("%Y-%m"))
    os.makedirs(subdir, exist_ok=True)
    stamp = now.strftime("%Y%m%dT%H%M%S")
    safe = os.path.basename(uploaded_file.name).replace(os.sep, "_")
    stored_path = os.path.join(subdir, f"{stamp}__{safe}")
    with open(stored_path, "wb") as out:
        for chunk in uploaded_file.chunks():
            out.write(chunk)
    return stored_path, safe


def _clean_header(value) -> str:
    if value is None:
        return ""
    s = str(value).replace("\r", " ").replace("\n", " ")
    return " ".join(s.split()).strip()


def _clean_cell(value):
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        try:
            return value.isoformat()
        except Exception:
            return str(value)
    if isinstance(value, (int, float, bool)):
        return value
    s = str(value).strip()
    return s if s else None


def _pick_sheet(wb):
    """Prefer a sheet named PO / Purchase Order, else the first sheet."""
    for want in PREFERRED_SHEET_NAMES:
        for name in wb.sheetnames:
            if name.strip().lower() == want.strip().lower():
                return wb[name], name
    return wb.worksheets[0], wb.sheetnames[0]


def _detect_header_row(ws, max_scan=15):
    """Return the 1-indexed row number that looks like the header row — the
    first row with at least 3 non-empty cells within the first ``max_scan``
    rows. Falls back to row 1 if nothing found."""
    for row_ix, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row_ix > max_scan:
            break
        non_blank = sum(1 for v in row if v is not None and str(v).strip())
        if non_blank >= 3:
            return row_ix
    return 1


def _parse_po_xlsx(stored_path):
    """Parse the PO xlsx and return
    {"po": {...}, "po_highlighted_columns": {...}, "filename": str, "sheet": str}.
    """
    wb = load_workbook(stored_path, data_only=True, read_only=True)
    ws, sheet_used = _pick_sheet(wb)

    header_row_ix = PO_HEADER_ROW_OVERRIDE or _detect_header_row(ws)

    # Read header + collect (col_index, cleaned_name) for real headers
    col_map = []
    rows_iter = ws.iter_rows(values_only=True)
    header_row = None
    for row_ix, row in enumerate(rows_iter, start=1):
        if row_ix == header_row_ix:
            header_row = row
            break
    if header_row is None:
        return {
            "filename": os.path.basename(stored_path),
            "sheet": sheet_used,
            "po": {"error": "header row not found", "row_count": 0, "headers": [], "rows": []},
            "po_highlighted_columns": {"row_count": 0, "rows": []},
        }

    for i, h in enumerate(header_row):
        name = _clean_header(h)
        if name:
            col_map.append((i, name))
    headers = [name for _i, name in col_map]

    all_rows = []
    for raw in rows_iter:
        row = {}
        any_val = False
        for i, name in col_map:
            v = _clean_cell(raw[i]) if i < len(raw) else None
            if v is None:
                continue
            row[name] = v
            any_val = True
        if any_val:
            all_rows.append(row)

    # Compute the subset for po_highlighted_columns
    if PO_HIGHLIGHTED_COLUMNS is None:
        subset_rows = all_rows                     # same as po tab
        subset_headers = headers
    elif not PO_HIGHLIGHTED_COLUMNS:
        subset_rows = []                           # explicitly empty
        subset_headers = []
    else:
        # keep only those columns in each row, preserving the order given
        subset_headers = list(PO_HIGHLIGHTED_COLUMNS)
        subset_rows = [
            {col: row[col] for col in PO_HIGHLIGHTED_COLUMNS if col in row}
            for row in all_rows
        ]

    return {
        "filename": os.path.basename(stored_path),
        "sheet": sheet_used,
        "po": {
            "row_count": len(all_rows),
            "headers": headers,
            "rows": all_rows,
        },
        "po_highlighted_columns": {
            "row_count": len(subset_rows),
            "headers": subset_headers,
            "rows": subset_rows,
        },
    }


def _forward_to_sheet(stored_path, original_filename, uploader_username, notes):
    url = _webhook_url()
    if not url:
        return {"attempted": False, "reason": "webhook not configured"}

    try:
        payload = _parse_po_xlsx(stored_path)
        payload["_meta"] = {
            "original_filename": original_filename,
            "uploader": uploader_username or "",
            "notes": notes or "",
        }
    except Exception as exc:
        log.exception("PO xlsx parse failed for %s", stored_path)
        return {"attempted": True, "ok": False, "error": f"parse failed: {exc}"}

    counts = {
        "po": payload["po"].get("row_count", 0),
        "po_highlighted_columns": payload["po_highlighted_columns"].get("row_count", 0),
    }

    try:
        resp = requests.post(url, json=payload, timeout=180)
        if 200 <= resp.status_code < 300:
            return {"attempted": True, "ok": True, "http_status": resp.status_code, "counts": counts}
        return {
            "attempted": True, "ok": False, "http_status": resp.status_code,
            "counts": counts, "body_preview": (resp.text or "")[:300],
        }
    except Exception as exc:
        log.exception("PO sheet-forward failed for %s", stored_path)
        return {"attempted": True, "ok": False, "error": str(exc), "counts": counts}


@api_view(["POST"])
@parser_classes([MultiPartParser])
def upload(request):
    uploaded = request.FILES.get("file")
    if not uploaded:
        return Response(
            {"detail": "No file attached under the 'file' key."},
            status=status.HTTP_400_BAD_REQUEST,
        )
    if not uploaded.name.lower().endswith((".xlsx", ".xlsm", ".xls", ".csv")):
        return Response(
            {"detail": "Please upload an .xlsx, .xls, or .csv file."},
            status=status.HTTP_400_BAD_REQUEST,
        )

    stored_path, original_name = _store_uploaded_file(uploaded)
    notes = ("" if request.data.get("notes") is None else str(request.data["notes"]))[:500]

    row = POUpload.objects.create(
        uploader=request.user,
        original_filename=original_name,
        stored_path=stored_path,
        notes=notes,
    )

    sheet = _forward_to_sheet(
        stored_path=stored_path,
        original_filename=original_name,
        uploader_username=request.user.username,
        notes=notes,
    )

    payload = POUploadSerializer(row).data
    payload["sheet_sync"] = sheet
    payload["target_sheet_url"] = TARGET_SHEET_URL
    return Response(payload, status=status.HTTP_201_CREATED)


@api_view(["GET"])
def list_uploads(request):
    qs = POUpload.objects.all()[:30]
    data = POUploadSerializer(qs, many=True).data
    return Response({
        "uploads": data,
        "target_sheet_url": TARGET_SHEET_URL,
        "webhook_configured": bool(_webhook_url()),
    })
