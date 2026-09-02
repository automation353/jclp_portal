"""REST endpoints for the PPC module.

  POST /api/ppc/uploads/new/  accept a file, save + record + forward to sheet
  GET  /api/ppc/uploads/      list the last 30 uploads

Django does the xlsx parsing (openpyxl, already available) and forwards a
clean JSON payload to the n8n webhook. n8n then writes rows to the target
Google Sheet — no xlsx library needed in the sandbox.

If JCLP_PPC_SHEET_WEBHOOK is blank the upload still succeeds; we just skip
the forward and report that in the response.
"""

import logging
import os

import requests
from django.utils import timezone
from openpyxl import load_workbook
from portal.notify import notify
from rest_framework import status
from rest_framework.decorators import api_view, parser_classes
from rest_framework.parsers import MultiPartParser
from rest_framework.response import Response

from .models import PPCUpload
from .serializers import PPCUploadSerializer


log = logging.getLogger(__name__)

UPLOAD_ROOT_DEFAULT = "/root/jclp_automation_portal/jcpl/uploads/ppc"

TARGET_SHEET_URL = (
    "https://docs.google.com/spreadsheets/d/"
    "1jg4EFvWaRjwJ4l8tt0FqBM0Fy3sHT-0y6LrWhYizMQ0/edit?gid=944387289"
)

# Fixed column schemas per sheet. We do NOT trust the source file's row-2
# header because it's incomplete on both CP Req and RM Req (missing labels
# for Item Code and UOM). Positions are 0-indexed column numbers.
#
# Format per sheet:
#   payload_key: the JSON key n8n reads from
#   data_start_row: 1-indexed row where DATA starts (rows above are ignored)
#   columns: list of (col_index, display_name)
SHEET_SCHEMAS = {
    "CP Req": {
        "payload_key": "cp_req",
        "data_start_row": 3,
        "columns": [
            (0, "BOM Code"),
            (1, "Item Code"),
            (2, "BOM Qty"),
            (3, "Qty in Nos"),
            (4, "Qty (Base UOM)"),
            (5, "UOM"),
            (6, "Part Description"),
        ],
    },
    "RM Req": {
        "payload_key": "rm_req",
        "data_start_row": 3,
        "columns": [
            (0, "BOM Code"),
            (1, "Item Code"),
            (2, "BOM Qty"),
            (3, "Qty in Nos"),
            (4, "Qty (Base UOM)"),
            (5, "UOM"),
            (6, "Part Description"),
        ],
    },
    "Quantity Sheet": {
        "payload_key": "quantity_sheet",
        "data_start_row": 4,        # row 2 = header, row 3 = blank spacer
        "columns": [
            (1, "Item Code"),
            (2, "Description"),
            (3, "UOM"),
            (4, "Qty (Base UOM)"),
            (5, "qty per kg"),
            (6, "Qty in nos (ERP)"),
            (7, "Reqd Qty"),
            (8, "Qty in nos (Correction)"),
        ],
    },
}


def _upload_root():
    return os.environ.get("JCLP_PPC_UPLOAD_ROOT", UPLOAD_ROOT_DEFAULT)


def _webhook_url():
    return os.environ.get("JCLP_PPC_SHEET_WEBHOOK", "").strip()


def _store_uploaded_file(uploaded_file) -> tuple:
    now = timezone.now()
    subdir = os.path.join(_upload_root(), now.strftime("%Y-%m"))
    os.makedirs(subdir, exist_ok=True)
    stamp = now.strftime("%Y%m%dT%H%M%S")
    safe = os.path.basename(uploaded_file.name).replace(os.sep, "_")
    stored_path = os.path.join(subdir, f"{stamp}__{safe}")
    with open(stored_path, "wb") as out:
        for chunk in uploaded_file.chunks():
            out.write(chunk)
    return stored_path, safe


def _clean_cell(value):
    """Return a JSON-safe scalar (or None) for a data cell."""
    if value is None:
        return None
    # datetime, date, time → ISO string. Everything else str-through JSON.
    if hasattr(value, "isoformat"):
        try:
            return value.isoformat()
        except Exception:
            return str(value)
    if isinstance(value, (int, float, bool)):
        return value
    s = str(value).strip()
    return s if s else None


def _extract_sheet(wb, sheet_name, schema):
    """Extract a sheet against a fixed column schema. Returns
    {'row_count', 'headers', 'rows'} or {'error': ...}.

    Fixed schemas make the extractor robust against the source file's
    incomplete header row on row 2 (both CP Req and RM Req have missing
    labels but the data underneath is well-shaped).
    """
    if sheet_name not in wb.sheetnames:
        return {"error": f"sheet {sheet_name!r} not found",
                "available_sheets": list(wb.sheetnames)}

    ws = wb[sheet_name]
    columns = schema["columns"]
    headers = [name for _i, name in columns]
    data_start_row = schema["data_start_row"]

    rows_iter = ws.iter_rows(values_only=True)
    # Advance past every row before data_start_row.
    for _ in range(data_start_row - 1):
        next(rows_iter, None)

    rows = []
    for raw in rows_iter:
        row = {}
        any_value = False
        for col_ix, name in columns:
            v = _clean_cell(raw[col_ix]) if col_ix < len(raw) else None
            if v is None:
                continue
            row[name] = v
            any_value = True
        if any_value:
            rows.append(row)

    return {"row_count": len(rows), "headers": headers, "rows": rows}


def _parse_ppc_xlsx(stored_path):
    """Parse the PPC xlsx per the SHEET_SCHEMAS map."""
    wb = load_workbook(stored_path, data_only=True, read_only=True)
    out = {"filename": os.path.basename(stored_path)}
    for sheet_name, schema in SHEET_SCHEMAS.items():
        out[schema["payload_key"]] = _extract_sheet(wb, sheet_name, schema)
    return out


def _forward_to_sheet(stored_path, original_filename, uploader_username, notes):
    url = _webhook_url()
    if not url:
        return {"attempted": False, "reason": "webhook not configured"}

    try:
        payload = _parse_ppc_xlsx(stored_path)
        payload["_meta"] = {
            "original_filename": original_filename,
            "uploader": uploader_username or "",
            "notes": notes or "",
        }
    except Exception as exc:
        log.exception("PPC xlsx parse failed for %s", stored_path)
        return {"attempted": True, "ok": False, "error": f"parse failed: {exc}"}

    counts = {
        schema["payload_key"]: (
            payload[schema["payload_key"]].get("row_count", 0)
            if isinstance(payload.get(schema["payload_key"]), dict) else 0
        )
        for schema in SHEET_SCHEMAS.values()
    }

    try:
        resp = requests.post(url, json=payload, timeout=180)
        if 200 <= resp.status_code < 300:
            return {
                "attempted": True, "ok": True,
                "http_status": resp.status_code,
                "counts": counts,
            }
        return {
            "attempted": True, "ok": False,
            "http_status": resp.status_code,
            "counts": counts,
            "body_preview": (resp.text or "")[:300],
        }
    except Exception as exc:
        log.exception("PPC sheet-forward failed for %s", stored_path)
        return {"attempted": True, "ok": False, "error": str(exc), "counts": counts}


@api_view(["POST"])
@parser_classes([MultiPartParser])
def upload(request):
    uploaded = request.FILES.get("file")
    if not uploaded:
        notify(
            "PPC upload rejected — no file",
            f"{request.user.get_username()} posted to the PPC upload endpoint "
            "with no file attached.",
        )
        return Response(
            {"detail": "No file attached under the 'file' key."},
            status=status.HTTP_400_BAD_REQUEST,
        )
    if not uploaded.name.lower().endswith((".xlsx", ".xlsm", ".xls", ".csv")):
        notify(
            "PPC upload rejected — wrong file type",
            f"{request.user.get_username()} tried to upload '{uploaded.name}', "
            "which isn't an .xlsx/.xls/.csv file.",
        )
        return Response(
            {"detail": "Please upload an .xlsx, .xls, or .csv file."},
            status=status.HTTP_400_BAD_REQUEST,
        )

    stored_path, original_name = _store_uploaded_file(uploaded)
    notes = ("" if request.data.get("notes") is None else str(request.data["notes"]))[:500]

    row = PPCUpload.objects.create(
        uploader=request.user,
        original_filename=original_name,
        stored_path=stored_path,
        notes=notes,
    )

    sheet = _forward_to_sheet(
        stored_path=stored_path,
        original_filename=original_name,
        uploader_username=request.user.username,
        notes=notes,
    )

    payload = PPCUploadSerializer(row).data
    payload["sheet_sync"] = sheet
    payload["target_sheet_url"] = TARGET_SHEET_URL

    if not sheet.get("attempted"):
        notify(
            f"PPC upload received — {original_name}",
            f"{request.user.get_username()} uploaded '{original_name}' "
            f"(upload #{row.id}). Sheet sync was not attempted "
            f"({sheet.get('reason')}) — the file is archived locally regardless.",
        )
    elif sheet.get("ok"):
        notify(
            f"PPC upload OK — {original_name}",
            f"{request.user.get_username()} uploaded '{original_name}' "
            f"(upload #{row.id}) and it synced to the sheet. Counts: "
            f"{sheet.get('counts')}",
        )
    else:
        notify(
            f"PPC upload received but SHEET SYNC FAILED — {original_name}",
            f"{request.user.get_username()} uploaded '{original_name}' "
            f"(upload #{row.id}) — the file is saved, but forwarding it to the "
            f"Google Sheet failed: {sheet.get('error') or sheet.get('body_preview')}",
        )
    return Response(payload, status=status.HTTP_201_CREATED)


@api_view(["GET"])
def list_uploads(request):
    qs = PPCUpload.objects.all()[:30]
    data = PPCUploadSerializer(qs, many=True).data
    return Response({
        "uploads": data,
        "target_sheet_url": TARGET_SHEET_URL,
        "webhook_configured": bool(_webhook_url()),
    })
