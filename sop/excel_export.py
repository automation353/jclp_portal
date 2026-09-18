"""Generate an Excel workbook from the S&OP Demand & Supply dashboard JSON."""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# ── Palette ─────────────────────────────────────────────────────────
_NAVY = "1A3A5C"
_RED = "B71C1C"
_OLIVE = "5D5A1E"
_TEAL = "2E5949"
_GREEN = "2E7D32"
_WHITE = "FFFFFF"
_LIGHT = "F5F8FF"
_ALT = "FAFBFC"

_hdr_font = Font(bold=True, color=_WHITE, size=11)
_hdr_fill = lambda hex_: PatternFill(start_color=hex_, end_color=hex_, fill_type="solid")
_bold = Font(bold=True, size=10)
_bold_green = Font(bold=True, color=_GREEN, size=10)
_num_fmt = '#,##0'
_num_fmt_d = '#,##0.0'
_pct_fmt = '0.0"%"'
_thin = Side(style="thin", color="D5DDE6")
_border = Border(bottom=_thin)


def _fmt(n):
    """Return numeric value or 0."""
    if n is None:
        return 0
    try:
        return float(n)
    except (TypeError, ValueError):
        return 0


def build_dashboard_excel(data):
    """Build and return an openpyxl Workbook from dashboard JSON (the dict
    returned by ``demand_supply_overview``)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Dashboard"
    ws.sheet_properties.pageSetUpPr = None  # default print

    # Column widths
    ws.column_dimensions["A"].width = 52
    for c in ("B", "C", "D", "E"):
        ws.column_dimensions[c].width = 18

    row = 1

    # ── Title ────────────────────────────────────────────────────────
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    c = ws.cell(row=row, column=1, value="Demand & Supply Visibility Dashboard")
    c.font = Font(bold=True, size=16, color=_NAVY)
    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    ws.cell(row=row, column=1, value="OEM Segment — Exported from JCPL Portal").font = Font(
        italic=True, color="666666", size=10
    )
    row += 2

    v = data.get("volume", {})
    h = data.get("health", {})

    # ── Helper: section header row ───────────────────────────────────
    def section(label, fill_hex, col_headers=None):
        nonlocal row
        for ci in range(1, 6):
            cell = ws.cell(row=row, column=ci)
            cell.fill = _hdr_fill(fill_hex)
            cell.font = _hdr_font
            cell.alignment = Alignment(horizontal="right" if ci > 1 else "left")
        ws.cell(row=row, column=1, value=label)
        if col_headers:
            for i, h_ in enumerate(col_headers):
                ws.cell(row=row, column=5 - len(col_headers) + 1 + i, value=h_)
        row += 1

    def data_row(label, t, f, r, bold_row=False, fmt=_num_fmt, label_font=None):
        nonlocal row
        ws.cell(row=row, column=1, value=label).font = label_font or (_bold if bold_row else Font(size=10))
        for ci, val in enumerate([t, f, r], start=3):
            c_ = ws.cell(row=row, column=ci, value=_fmt(val))
            c_.number_format = fmt
            c_.font = _bold if bold_row else Font(size=10)
            c_.alignment = Alignment(horizontal="right")
            c_.border = _border
        ws.cell(row=row, column=1).border = _border
        row += 1

    # ── VOLUME & FULFILMENT ──────────────────────────────────────────
    section("VOLUME & FULFILMENT", _NAVY, ["Total", "Focus", "Regular"])
    data_row("Total Items", v.get("total_items"), v.get("focus_items"), v.get("regular_items"), bold_row=True)

    vol_rows = [
        ("Original Forecast (units)", "forecast"),
        ("Ops+Sales-Reviewed Forecast Committed", "committed"),
        ("Actual Sales Demand", "actual_sales"),
        ("Additional Demand (Punched, Not Committed)", "additional_demand"),
        ("Actual Production (units)", "production"),
        ("Actual Dispatch (units)", "dispatch"),
        ("Closing Inventory (units)", "closing_inventory"),
        ("Free Inventory above Green Level", "free_inventory"),
    ]
    for label, key in vol_rows:
        data_row(
            label,
            v.get("total", {}).get(key),
            v.get("focus", {}).get(key),
            v.get("regular", {}).get(key),
            label_font=_bold_green,
        )
    row += 1

    # ── HEALTH RATIOS ────────────────────────────────────────────────
    section("HEALTH RATIOS", _NAVY, ["Total", "Focus", "Regular"])
    health_rows = [
        ("Dispatch Fulfilment (Dispatch ÷ Committed)", "dispatch_fulfilment"),
        ("Forecast Accuracy (Hit Rate — within 10%)", "forecast_accuracy"),
        ("Production Coverage (capped)", "production_coverage"),
    ]
    for label, key in health_rows:
        t_val = h.get(key, h.get("total", {}).get(key, 0))
        f_val = h.get("focus", {}).get(key, 0)
        r_val = h.get("regular", {}).get(key, 0)
        nonlocal_row = row
        ws.cell(row=row, column=1, value=label).font = Font(size=10, color=_GREEN)
        ws.cell(row=row, column=1).border = _border
        for ci, val in enumerate([t_val, f_val, r_val], start=3):
            c_ = ws.cell(row=row, column=ci, value=f"{_fmt(val)}%")
            c_.font = Font(bold=True, size=10, color=_GREEN if _fmt(val) >= 70 else "C62828")
            c_.alignment = Alignment(horizontal="right")
            c_.border = _border
        row += 1
    row += 1

    # ── DEMAND VALUE (₹) ─────────────────────────────────────────────
    dv = data.get("demand_value")
    if dv:
        section("DEMAND VALUE (₹) — ASP Weighted", _NAVY, ["Total", "Focus", "Regular"])
        for label, key in [
            ("Original Forecast Value", "original_forecast_value"),
            ("Committed Demand Value", "committed_demand_value"),
        ]:
            data_row(
                label,
                dv.get("total", {}).get(key),
                dv.get("focus", {}).get(key),
                dv.get("regular", {}).get(key),
                label_font=_bold_green,
            )
        row += 1

    # ── EXCEPTIONS — ITEM COUNTS ─────────────────────────────────────
    ex = data.get("exceptions", {})
    section("EXCEPTIONS — ITEM COUNTS", _RED, ["Count", "Focus", "Regular"])
    exc_rows = [
        ("Items with Production-Driven Shortfall", "production_shortfall"),
        ("Items in Production Surplus", "production_surplus"),
        ("Items with Excess Opening Stock", "excess_opening_stock"),
        ("Items with Dispatch Gap", "dispatch_gap"),
        ("Items Below Green Level", "below_green_level"),
        ("Items At/Above Green Level", "above_green_level"),
    ]
    for label, key in exc_rows:
        d = ex.get(key, {})
        if isinstance(d, dict):
            data_row(label, d.get("total", 0), d.get("focus", 0), d.get("regular", 0))
        else:
            data_row(label, d, 0, 0)
    row += 1

    # ── TYPE BREAKDOWN ───────────────────────────────────────────────
    types = data.get("type_breakdown", [])
    if types:
        section("TYPE BREAKDOWN — MTS / MTO / TBC", _OLIVE)
        # Header
        for ci, h_ in enumerate(["Type", "Items", "Share %", "Forecast", "Production"], 1):
            c_ = ws.cell(row=row, column=ci, value=h_)
            c_.font = _bold
            c_.fill = PatternFill(start_color="F5F5F0", end_color="F5F5F0", fill_type="solid")
            c_.border = _border
        row += 1
        type_labels = {"MTS": "MTS — Make to Stock", "MTO": "MTO — Make to Order", "TBC": "TBC — Classification Pending"}
        for t_ in types:
            ws.cell(row=row, column=1, value=type_labels.get(t_.get("type", ""), t_.get("type", ""))).font = Font(size=10, bold=True)
            ws.cell(row=row, column=2, value=_fmt(t_.get("items"))).number_format = _num_fmt
            ws.cell(row=row, column=3, value=f'{t_.get("share_pct", 0)}%')
            ws.cell(row=row, column=4, value=_fmt(t_.get("forecast"))).number_format = _num_fmt
            ws.cell(row=row, column=5, value=_fmt(t_.get("production"))).number_format = _num_fmt
            for ci in range(1, 6):
                ws.cell(row=row, column=ci).border = _border
                if ci > 1:
                    ws.cell(row=row, column=ci).alignment = Alignment(horizontal="right")
            row += 1
        row += 1

    # ── SALES INSIGHT TAGS ───────────────────────────────────────────
    ins = data.get("insights_summary", {})
    total_items = v.get("total_items", 1) or 1
    for section_title, tag_key, fill_hex, order in [
        ("SALES INSIGHT TAGS", "sales_tags", _TEAL,
         ["On Target", "Over-Delivered", "Short Delivery", "Dispatch Gap",
          "Production Shortfall", "Forecast Bullseye", "Forecast Drift (Minor)",
          "Forecast Drift (Moderate)", "Forecast Drift (Major)"]),
        ("OPERATIONS INSIGHT TAGS", "ops_tags", _OLIVE,
         ["Healthy Stock", "Production Shortfall", "Below Safety", "Stock-Out Risk",
          "Excess Inventory", "Production Surplus", "Unsold MTO Stock", "Over-Production",
          "Type TBC"]),
    ]:
        tags = ins.get(tag_key, {})
        if tags:
            section(section_title, fill_hex)
            for ci, h_ in enumerate(["Tag", "Items", "Share %"], 1):
                c_ = ws.cell(row=row, column=ci, value=h_)
                c_.font = _bold
                c_.fill = PatternFill(start_color="F5F5F0", end_color="F5F5F0", fill_type="solid")
                c_.border = _border
            row += 1
            for tag in order:
                count = tags.get(tag, 0)
                ws.cell(row=row, column=1, value=tag).border = _border
                c2 = ws.cell(row=row, column=2, value=_fmt(count))
                c2.number_format = _num_fmt
                c2.alignment = Alignment(horizontal="right")
                c2.border = _border
                pct = round(count / total_items * 100, 1) if total_items else 0
                c3 = ws.cell(row=row, column=3, value=f"{pct}%")
                c3.alignment = Alignment(horizontal="right")
                c3.border = _border
                row += 1
            row += 1

    # ── TOP 10 sections ──────────────────────────────────────────────
    top10_sections = [
        ("TOP 10 — PRODUCTION-DRIVEN SHORTFALL", "top10_shortfall",
         _RED, ["committed", "production", "dispatch", "shortfall"]),
        ("TOP 10 — DISPATCH GAP", "top10_dispatch_gap",
         _RED, ["committed", "dispatch", "dispatch_pct", "gap"]),
        ("TOP 10 — EXCESS FREE INVENTORY", "top10_free_inventory",
         _OLIVE, ["closing", "green_level", "free_inv", "x_green"]),
        ("TOP 10 — BELOW GREEN LEVEL", "top10_below_green",
         _RED, ["closing", "green_level", "deficit"]),
    ]
    for title, key, fill_hex, cols in top10_sections:
        items = data.get(key, [])
        if not items:
            continue
        section(title, fill_hex)
        # Sub-header
        sub_headers = ["Rank", "Item Code", "Type"] + [c.replace("_", " ").title() for c in cols]
        for ci, h_ in enumerate(sub_headers, 1):
            c_ = ws.cell(row=row, column=ci, value=h_)
            c_.font = _bold
            c_.fill = PatternFill(start_color="F5F5F0", end_color="F5F5F0", fill_type="solid")
            c_.border = _border
        row += 1
        for rank, item in enumerate(items, 1):
            ws.cell(row=row, column=1, value=rank).alignment = Alignment(horizontal="center")
            ws.cell(row=row, column=2, value=item.get("item_code", ""))
            ws.cell(row=row, column=3, value=item.get("type", ""))
            for ci, col in enumerate(cols, 4):
                val = item.get(col, 0)
                c_ = ws.cell(row=row, column=ci, value=_fmt(val) if isinstance(val, (int, float)) or val is None else val)
                c_.alignment = Alignment(horizontal="right")
                if isinstance(val, (int, float)):
                    c_.number_format = _num_fmt_d
            for ci in range(1, 4 + len(cols)):
                ws.cell(row=row, column=ci).border = _border
            row += 1
        row += 1

    # ── EXCEPTIONS — Qty & ₹ VALUE ───────────────────────────────────
    ev = data.get("exceptions_value")
    if ev:
        section("EXCEPTIONS — Qty & ₹ VALUE", _RED)
        # Header rows for TOTAL / MTS / MTO / TBC
        merge_headers = ["", "TOTAL", "", "MTS", "", "MTO", "", "TBC", ""]
        sub_headers = ["Bucket", "Qty", "₹ Value", "Qty", "₹ Value", "Qty", "₹ Value", "Qty", "₹ Value"]
        # Adjust col widths for this wider section
        for ci, h_ in enumerate(sub_headers, 1):
            c_ = ws.cell(row=row, column=ci, value=h_)
            c_.font = _bold
            c_.fill = PatternFill(start_color="F5F5F0", end_color="F5F5F0", fill_type="solid")
            c_.border = _border
            if ci > 1:
                c_.alignment = Alignment(horizontal="right")
                ws.column_dimensions[get_column_letter(ci)].width = 14
        row += 1

        ev_rows = [
            ("Production Surplus", "production_surplus"),
            ("Excess Opening Stock", "excess_opening"),
            ("Demand Reduction Adjustment", "demand_reduction"),
            ("Excess Dispatched", "excess_dispatched"),
            ("Uncovered Shortfall", "uncovered_shortfall"),
        ]
        for label, key in ev_rows:
            d = ev.get(key, {})
            if not d:
                continue
            ws.cell(row=row, column=1, value=label).font = Font(bold=True, size=10)
            ws.cell(row=row, column=1).border = _border
            type_keys = ["total", "MTS", "MTO", "TBC"]
            ci = 2
            for tk in type_keys:
                td = d.get(tk, {})
                c_q = ws.cell(row=row, column=ci, value=_fmt(td.get("qty", 0)))
                c_q.number_format = _num_fmt_d
                c_q.alignment = Alignment(horizontal="right")
                c_q.border = _border
                ci += 1
                c_v = ws.cell(row=row, column=ci, value=_fmt(td.get("value", 0)))
                c_v.number_format = _num_fmt
                c_v.alignment = Alignment(horizontal="right")
                c_v.border = _border
                ci += 1
            row += 1

    return wb


def dashboard_to_bytes(data):
    """Return the Excel workbook as bytes (ready for HttpResponse)."""
    wb = build_dashboard_excel(data)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
