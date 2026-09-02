"""Fast xlsx parser for the MTO_MTS workbook.

Follows Section 3 of the automation brief:

  - Sheet must be named ``MTO_MTS`` (case/whitespace tolerant).
  - Anchor on the row containing both "Item Group" and "Item Code" — extra
    label rows above the data are allowed.
  - Every column strictly between "Item Code" and (if present) "Status" is a
    month column. Empty right-hand columns are future months, kept for
    reference but marked blank.
  - Every value must be "MTS", "MTO", or blank. Anything else is recorded
    as an exception but does not abort the whole run.
  - Excel dates in the header row are accepted — they are formatted as
    ``"MMM YYYY"`` so both the DB and the UI stay readable.

**Performance-critical note.** openpyxl's ``read_only=True`` mode is a
streaming reader — ``sheet.cell(row=X, column=Y)`` is O(row) per call
because it iterates the underlying XML from the start every time. On a
1600-row file that turns the parse into O(n²) and pins a CPU for many
minutes. We therefore do a single pass via ``iter_rows(values_only=True)``.

Raises ``MtoMtsParseError`` when the file's structure is unrecoverable —
matching the brief's "stop and raise an alert instead of guessing" rule.
"""

from dataclasses import dataclass, field
from datetime import date, datetime
from typing import List

from openpyxl import load_workbook


SHEET_NAME_CANDIDATES = ("MTO_MTS", "MTO/MTS", "MTOMTS")
VALID_VALUES = {"MTS", "MTO", ""}
# The V2 template has a segment-legend block up top, so the data header row
# can be as far down as row 10 — scan generously.
MAX_HEADER_SCAN_ROWS = 40


class MtoMtsParseError(Exception):
    """Structural problem with the file — halt, don't guess."""


@dataclass
class ParsedCell:
    month_label: str
    month_index: int
    value: str


@dataclass
class ParsedItem:
    item_group: str
    item_code: str
    segment: str = ""              # optional "Segment" column (Change 15)
    account_description: str = ""  # optional "Item A/C Description" column
    cells: List[ParsedCell] = field(default_factory=list)


@dataclass
class ParseResult:
    items: List[ParsedItem] = field(default_factory=list)
    month_labels: List[str] = field(default_factory=list)
    exceptions: List[str] = field(default_factory=list)


def _norm(value) -> str:
    """Coerce a workbook cell into a stripped display string. Excel dates
    become 'MMM YYYY' so month-column headers stay human-readable."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%b %Y")
    if isinstance(value, date):
        return value.strftime("%b %Y")
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def _canonical_sheet_key(name: str) -> str:
    return "".join(ch for ch in name.upper() if ch.isalnum())


def _find_sheet(workbook):
    wanted = {_canonical_sheet_key(c) for c in SHEET_NAME_CANDIDATES}
    for sheet in workbook.worksheets:
        if _canonical_sheet_key(sheet.title) in wanted:
            return sheet
    return None


def parse_workbook(file_path: str) -> ParseResult:
    workbook = load_workbook(file_path, data_only=True, read_only=True)

    sheet = _find_sheet(workbook)
    if sheet is None:
        raise MtoMtsParseError(
            f"Workbook has no 'MTO_MTS' sheet (found: {workbook.sheetnames})."
        )

    # ONE streaming pass. Anything requiring random-access on the sheet
    # (sheet.cell(...)) would re-iterate from row 1 each call and blow up.
    row_iter = sheet.iter_rows(values_only=True)

    # --- Locate the header row ------------------------------------------
    header_row_ix = None
    header_values: List[str] = []
    for row_ix, raw_row in enumerate(row_iter, start=1):
        if row_ix > MAX_HEADER_SCAN_ROWS:
            break
        cells = [_norm(v) for v in raw_row]
        lower = {c.lower() for c in cells}
        if "item group" in lower and "item code" in lower:
            header_row_ix = row_ix
            header_values = cells
            break

    if header_row_ix is None:
        raise MtoMtsParseError(
            "Could not find a header row containing both 'Item Group' and "
            f"'Item Code' within the first {MAX_HEADER_SCAN_ROWS} rows."
        )

    # --- Locate the columns from the header row --------------------------
    group_col = code_col = status_col = segment_col = ac_col = None
    month_cols: List[tuple] = []      # (0-based index, label)
    for ix, label in enumerate(header_values):
        low = label.lower()
        if low == "item group":
            group_col = ix
        elif low == "item code":
            code_col = ix
        elif low == "status":
            status_col = ix
        elif low in (
            "revised segment",     # V2 template label — 4 broad categories
            "segment",
            "category",
            "business segment",
        ):
            # The high-level category. Feeds the Category dropdown in the UI.
            segment_col = ix
        elif low in (
            "party a/c description",   # V2 template label
            "item a/c description",    # legacy V1 label
            "item ac description",
            "party ac description",
            "account description",
            "a/c description",
        ):
            # Party/account description — the granular account code (still
            # stored so table + future views can show it).
            ac_col = ix
        elif label:
            month_cols.append((ix, label))

    if group_col is None or code_col is None:
        raise MtoMtsParseError(
            "Header row is missing 'Item Group' and/or 'Item Code'."
        )

    # Only columns strictly to the right of Item Code and strictly to the
    # left of Status (if present) count as month columns. The Segment and
    # A/C Description columns, if present, are also excluded.
    filtered: List[tuple] = []
    for ix, label in month_cols:
        if ix <= code_col:
            continue
        if status_col is not None and ix >= status_col:
            continue
        if segment_col is not None and ix == segment_col:
            continue
        if ac_col is not None and ix == ac_col:
            continue
        filtered.append((ix, label))
    filtered.sort(key=lambda pair: pair[0])
    month_labels = [label for _ix, label in filtered]

    result = ParseResult(month_labels=month_labels)

    # --- Read data rows in the same stream --------------------------------
    for offset, raw_row in enumerate(row_iter, start=1):
        item_group = _norm(raw_row[group_col]) if group_col < len(raw_row) else ""
        item_code = _norm(raw_row[code_col]) if code_col < len(raw_row) else ""

        if not item_group and not item_code:
            continue

        segment = ""
        if segment_col is not None and segment_col < len(raw_row):
            segment = _norm(raw_row[segment_col])

        account_description = ""
        if ac_col is not None and ac_col < len(raw_row):
            account_description = _norm(raw_row[ac_col])

        item = ParsedItem(
            item_group=item_group,
            item_code=item_code,
            segment=segment,
            account_description=account_description,
        )
        for month_index, (ix, label) in enumerate(filtered):
            raw = _norm(raw_row[ix]) if ix < len(raw_row) else ""
            up = raw.upper()
            if up in VALID_VALUES:
                item.cells.append(
                    ParsedCell(month_label=label, month_index=month_index, value=up)
                )
            else:
                data_row_number = header_row_ix + offset
                result.exceptions.append(
                    f"Row {data_row_number} ({item_code}): unexpected value "
                    f"'{raw}' in month '{label}' — treated as blank."
                )
                item.cells.append(
                    ParsedCell(month_label=label, month_index=month_index, value="")
                )

        result.items.append(item)

    if not result.items:
        raise MtoMtsParseError("No data rows found under the header.")

    return result
