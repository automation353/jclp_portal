"""Parser for EBQ data → batch_ebq table (W1.8).

Supports two file formats:
  1. Monitoring.xlsx — sheets "EBQ"/"Batch Qty", family-level EBQ
  2. EBQ Qualification.xlsx — "Data" sheet, per-item EBQ with ERP Code
Level: L0
"""

import logging

from openpyxl import load_workbook

from ..field_maps.batch_ebq import FIELD_MAP, FIELD_MAP_QUAL, TABLE_KEY
from .base import detect_header_row, extract_rows

log = logging.getLogger(__name__)

TABLE_KEY = TABLE_KEY  # noqa: F841


def _parse_qualification_sheet(file_path, sheet_name):
    """Parse per-item EBQ from EBQ Qualification.xlsx Data sheet."""
    wb = load_workbook(file_path, read_only=True, data_only=True)
    ws = wb[sheet_name]

    header_map = {v.upper(): k for k, v in FIELD_MAP_QUAL.items()}
    header_row_idx = None
    col_map = {}

    for i, row in enumerate(ws.iter_rows(values_only=True)):
        cells = [str(c).strip() if c is not None else "" for c in row]
        matched = 0
        trial_map = {}
        for j, cell in enumerate(cells):
            key = header_map.get(cell.upper())
            if key and key not in trial_map.values():
                trial_map[j] = key
                matched += 1
        if matched >= 3:
            header_row_idx = i
            col_map = trial_map
            break

    if header_row_idx is None:
        wb.close()
        return []

    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i <= header_row_idx:
            continue
        vals = list(row)
        erp_code = None
        ebq_val = None
        record = {}
        for j, key in col_map.items():
            if j < len(vals):
                record[key] = vals[j]
                if key == "erp_code":
                    erp_code = str(vals[j]).strip() if vals[j] is not None else ""
                if key == "ebq_qty":
                    try:
                        ebq_val = float(vals[j]) if vals[j] is not None else 0
                    except (ValueError, TypeError):
                        ebq_val = 0
        if erp_code and ebq_val and ebq_val > 0:
            rows.append(record)

    wb.close()
    log.info("Parsed %d per-item EBQ rows from sheet %r", len(rows), sheet_name)
    return rows


def parse(file_path, sheet_name=None):
    """Parse EBQ / batch data from Monitoring.xlsx or EBQ Qualification.xlsx."""
    wb = load_workbook(file_path, read_only=True, data_only=True)
    sheets = wb.sheetnames
    wb.close()

    # Try EBQ Qualification format first (has "Data" sheet)
    for name in sheets:
        if name.lower() == "data":
            rows = _parse_qualification_sheet(file_path, name)
            if rows:
                log.info("Parsed %d total EBQ rows (qualification format) from %s",
                         len(rows), file_path)
                return rows

    # Fall back to old Monitoring.xlsx format
    all_rows = []
    for name in sheets:
        nl = name.lower()
        if "ebq" in nl or "batch" in nl:
            try:
                rows = extract_rows(file_path, name, FIELD_MAP)
                all_rows.extend(rows)
                log.info("Parsed %d rows from sheet %r", len(rows), name)
            except ValueError as e:
                log.info("Skipping sheet %r: %s", name, e)

    if not all_rows:
        all_rows = extract_rows(file_path, None, FIELD_MAP)

    log.info("Parsed %d total EBQ/batch rows from %s", len(all_rows), file_path)
    return all_rows
