"""Parser for MpsSS.xlsm — Demand Data + Dispatch Data (L3 history).

Source: MpsSS.xlsm → "Demand Data" and "Dispatch Data" sheets.
~3,000 rows. Monthly demand and dispatch history with rate, value.
"""

import logging

from openpyxl import load_workbook

from ..field_maps.mps_history import FIELD_MAP
from .base import clean_cell, detect_header_row

TABLE_KEY = "mps_history"

log = logging.getLogger(__name__)

_DEMAND_SHEETS = ["Demand Data", "Demand"]
_DISPATCH_SHEETS = ["Dispatch Data", "Dispatch"]


def _find_sheet(wb, candidates):
    for sn in wb.sheetnames:
        sn_lower = sn.lower().strip()
        for cand in candidates:
            if cand.lower() in sn_lower:
                return sn
    return None


def _extract(wb, sheet_name, header_to_key, known_headers):
    """Extract rows from one sheet."""
    if sheet_name is None:
        return []
    ws = wb[sheet_name]
    try:
        header_row_idx, col_map = detect_header_row(ws, known_headers)
    except ValueError:
        log.warning("Sheet %s — no matching headers", sheet_name)
        return []

    col_to_key = {}
    for col_idx, header_str in col_map.items():
        key = header_to_key.get(header_str.strip())
        if key:
            col_to_key[col_idx] = key

    rows = []
    for row_idx, raw_row in enumerate(ws.iter_rows(values_only=True)):
        if row_idx <= header_row_idx:
            continue
        record = {}
        any_value = False
        for col_idx, key in col_to_key.items():
            v = clean_cell(raw_row[col_idx]) if col_idx < len(raw_row) else None
            if v is not None:
                record[key] = v
                any_value = True
        if any_value and record.get("item_code"):
            rows.append(record)
    return rows


def parse(file_path):
    """Parse MPS history from MpsSS.xlsm.
    Merges Demand Data + Dispatch Data sheets.
    """
    wb = load_workbook(file_path, data_only=True, read_only=True)

    header_to_key = {v.strip(): k for k, v in FIELD_MAP.items()}
    known_headers = list(header_to_key.keys())

    all_rows = []

    demand_sn = _find_sheet(wb, _DEMAND_SHEETS)
    all_rows.extend(_extract(wb, demand_sn, header_to_key, known_headers))

    dispatch_sn = _find_sheet(wb, _DISPATCH_SHEETS)
    all_rows.extend(_extract(wb, dispatch_sn, header_to_key, known_headers))

    wb.close()
    log.info("Parsed %d MPS history rows from %s", len(all_rows), file_path)
    return all_rows
