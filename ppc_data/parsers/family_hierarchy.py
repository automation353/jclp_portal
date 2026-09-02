"""Parser for Family Hierarchy table (W1.2).

Sources: Product Group Mapping.xlsx + Monitoring.xlsx!Family Group
Both are merged into one table. If only one file is uploaded, it
parses what it can — the merge happens when both have been loaded.

This parser handles either file. The caller (or a management command)
is responsible for merging when both are present.
Level: L0
"""

import logging

from ..field_maps.family_hierarchy import FIELD_MAP, TABLE_KEY
from .base import extract_rows

log = logging.getLogger(__name__)

TABLE_KEY = TABLE_KEY  # noqa: F841


def parse(file_path, sheet_name=None):
    """Parse a family hierarchy source file.

    Auto-detects: if the file contains a "Family Group" sheet, use it.
    Otherwise, try the first sheet (Product Group Mapping style).
    """
    from openpyxl import load_workbook

    wb = load_workbook(file_path, read_only=True, data_only=True)
    sheets = wb.sheetnames
    wb.close()

    # Prefer "Family Group" sheet from Monitoring.xlsx
    target = None
    for name in sheets:
        if "family" in name.lower() and "group" in name.lower():
            target = name
            break

    rows = extract_rows(file_path, target, FIELD_MAP)
    log.info(
        "Parsed %d family hierarchy rows from %s (sheet: %s)",
        len(rows), file_path, target or "first",
    )
    return rows
