"""Parser for ASP for FG.xlsx → rate_asp table (W1.16).

Sheet: "Sheet1"
Rows: ~1,837
Simple key-value load.
Level: L0
"""

import logging

from openpyxl import load_workbook

from ..field_maps.rate_asp import FIELD_MAP, TABLE_KEY
from .base import extract_rows

log = logging.getLogger(__name__)

TABLE_KEY = TABLE_KEY  # noqa: F841


def parse(file_path, sheet_name=None):
    """Parse ASP / rate file. Prefers "Sheet1", falls back to first."""
    if not sheet_name:
        wb = load_workbook(file_path, read_only=True, data_only=True)
        sheets = wb.sheetnames
        wb.close()
        sheet_name = "Sheet1" if "Sheet1" in sheets else None

    rows = extract_rows(file_path, sheet_name, FIELD_MAP)
    log.info("Parsed %d ASP rows from %s", len(rows), file_path)
    return rows
