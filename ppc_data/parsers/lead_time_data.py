"""Parser for Lead Time Data.xlsx → lead_time table (W1.9).

Sheet: "Sheet2" (1,684 rows)
WARNING: Last updated May 2023 — must be revalidated.
Level: L0
"""

import logging

from openpyxl import load_workbook

from ..field_maps.lead_time import FIELD_MAP, TABLE_KEY
from .base import extract_rows

log = logging.getLogger(__name__)

TABLE_KEY = TABLE_KEY  # noqa: F841


def parse(file_path, sheet_name=None):
    """Parse lead time data. Prefers "Sheet2", falls back to first."""
    if not sheet_name:
        wb = load_workbook(file_path, read_only=True, data_only=True)
        sheets = wb.sheetnames
        wb.close()
        sheet_name = "Sheet2" if "Sheet2" in sheets else None

    rows = extract_rows(file_path, sheet_name, FIELD_MAP)
    log.info("Parsed %d lead time rows from %s", len(rows), file_path)
    return rows
