"""Parser for FG.xlsx — Stock Statement Valuation Report (all sites).

Source: ERP FG Stock/FG.xlsx
Table key: fg_stock_statement
Provides: Opening Balance, FG (closing qty) for R3SS
"""

import logging

from openpyxl import load_workbook

log = logging.getLogger(__name__)

TABLE_KEY = "fg_stock_statement"

_SHEET_CANDIDATES = [
    "Stock Statment Valuation Report",
    "Stock Statement Valuation Report",
    "Stock Statement",
    "Sheet1",
]

_HEADER_RULES = [
    ("ITEM NAME", "erp_code"),
    ("ITEM CODE", "item_code_erp"),
    ("OPENING", "opening_qty"),
    ("RECEIPT", "receipts"),
    ("ISSUE", "issues"),
    ("CLOSING", "closing_qty"),
]


def _find_sheet(wb):
    for cand in _SHEET_CANDIDATES:
        for sn in wb.sheetnames:
            if cand.lower() == sn.lower().strip():
                return sn
    return wb.sheetnames[0]


def _num(v):
    if v is None:
        return 0.0
    try:
        return float(v)
    except (ValueError, TypeError):
        return 0.0


def parse(file_path):
    wb = load_workbook(file_path, data_only=True)
    ws = wb[_find_sheet(wb)]

    header_row_idx = None
    col_map = {}

    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i >= 10:
            break
        cells = [str(c).strip().upper() if c is not None else "" for c in row]
        matched = 0
        trial = {}
        for j, cell in enumerate(cells):
            for prefix, key in _HEADER_RULES:
                if cell.startswith(prefix) and key not in trial.values():
                    trial[j] = key
                    matched += 1
                    break
        if matched >= 3:
            header_row_idx = i
            col_map = trial
            break

    if header_row_idx is None:
        wb.close()
        log.warning("No matching headers found in %s", file_path)
        return []

    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i <= header_row_idx:
            continue
        vals = list(row)
        erp = None
        record = {}
        for j, key in col_map.items():
            if j >= len(vals):
                continue
            v = vals[j]
            if key == "erp_code":
                erp = str(v).strip() if v is not None else ""
                record[key] = erp
            elif key == "item_code_erp":
                record[key] = str(v).strip() if v is not None else ""
            else:
                record[key] = _num(v)

        if erp:
            rows.append(record)

    wb.close()
    log.info("Parsed %d rows from fg_stock_statement %s", len(rows), file_path)
    return rows
