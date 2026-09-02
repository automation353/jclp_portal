"""Generic passthrough parser for S&OP Excel files.

Reads any .xlsx file — auto-detects the header row (first row with 3+
non-empty cells), normalises header names to snake_case keys, and returns
all data rows as dicts.

Handles TCS iON ERP exports which:
  - Have multiple sheets (picks the largest data sheet, not the title page)
  - Need read_only=False for openpyxl to see all columns

Used for all 5 S&OP upload types:
  - sop_dpr            (Stock Ledger / DPR)
  - sop_forecast       (Forecast vs Sales Report)
  - sop_opening_stock  (Stock Statement Valuation Report)
  - sop_green_level    (PPC Green Level Quantities)
  - sop_sales_register (Sales Invoice Register)
"""

import logging
import re

from openpyxl import load_workbook

from .base import clean_cell

log = logging.getLogger(__name__)

# ── Table keys for S&OP uploads ──────────────────────────────────────
SOP_TABLE_KEYS = [
    "sop_dpr",
    "sop_forecast",
    "sop_opening_stock",
    "sop_green_level",
    "sop_sales_register",
]


def _slugify(header):
    """Convert an Excel header to a snake_case key.

    'Item Code'      → 'item_code'
    'Qty (MT)'       → 'qty_mt'
    'Opening Bal.'   → 'opening_bal'
    'Date / Month'   → 'date_month'
    """
    s = str(header).strip()
    s = re.sub(r"[^\w\s]", " ", s)      # strip punctuation
    s = re.sub(r"\s+", "_", s.strip())   # spaces → underscores
    return s.lower()[:60]                # max 60 chars


def _detect_header_row(ws, max_rows=25, min_cols=3):
    """Find the first row with at least ``min_cols`` non-empty cells.

    Returns (row_index_0based, [header_strings]).
    """
    for row_idx, row in enumerate(ws.iter_rows(max_row=max_rows, values_only=True)):
        non_empty = [c for c in row if c is not None and str(c).strip()]
        if len(non_empty) >= min_cols:
            return row_idx, list(row)
    raise ValueError(
        f"No header row found in the first {max_rows} rows "
        f"(need at least {min_cols} non-empty cells)."
    )


def _pick_data_sheet(wb):
    """Pick the best data sheet from a workbook.

    TCS iON exports have: 'main report' (title page), 'Search Criteria'
    (empty/metadata), and one or more data sheets. We pick the sheet with
    the most columns in its first few rows — that's the actual data.
    """
    if len(wb.sheetnames) == 1:
        return wb.active

    best_sheet = None
    best_score = 0

    for name in wb.sheetnames:
        # Skip common metadata/title sheets
        lower = name.lower()
        if lower in ("search criteria", "search_criteria"):
            continue

        ws = wb[name]
        # Score = max non-empty cells in any of the first 25 rows
        score = 0
        for row in ws.iter_rows(max_row=25, values_only=True):
            non_empty = sum(1 for c in row if c is not None and str(c).strip())
            score = max(score, non_empty)

        if score > best_score:
            best_score = score
            best_sheet = ws

    # Fall back to active if nothing found
    return best_sheet or wb.active


def parse(file_path, sheet_name=None, table_key=None):
    """Parse any S&OP Excel file generically.

    Returns list of dicts keyed by slugified header names.

    Uses read_only=False to handle TCS iON exports where read_only mode
    doesn't detect all columns properly.
    """
    # read_only=False is needed for TCS iON exports
    wb = load_workbook(file_path, data_only=True, read_only=False)

    if sheet_name:
        if sheet_name not in wb.sheetnames:
            available = ", ".join(wb.sheetnames)
            raise ValueError(
                f"Sheet {sheet_name!r} not found. Available: {available}"
            )
        ws = wb[sheet_name]
    else:
        ws = _pick_data_sheet(wb)

    log.info(
        "SOP parser: using sheet %r from %s (available: %s)",
        ws.title, file_path, wb.sheetnames,
    )

    header_row_idx, raw_headers = _detect_header_row(ws)

    # Build column index → key mapping (skip empty headers)
    col_keys = {}
    seen = set()
    for col_idx, hdr in enumerate(raw_headers):
        if hdr is None or not str(hdr).strip():
            continue
        key = _slugify(hdr)
        if not key:
            continue
        # Deduplicate: append _2, _3, … on collisions
        base_key = key
        n = 2
        while key in seen:
            key = f"{base_key}_{n}"
            n += 1
        seen.add(key)
        col_keys[col_idx] = key

    if not col_keys:
        raise ValueError("No usable headers found in the detected header row.")

    # Extract data rows
    rows = []
    for row_idx, raw_row in enumerate(ws.iter_rows(values_only=True)):
        if row_idx <= header_row_idx:
            continue
        record = {}
        any_value = False
        for col_idx, key in col_keys.items():
            v = clean_cell(raw_row[col_idx]) if col_idx < len(raw_row) else None
            if v is not None:
                record[key] = v
                any_value = True
        if any_value:
            rows.append(record)

    wb.close()

    log.info(
        "SOP generic parser: %d rows, %d columns from %s (sheet=%s, table_key=%s)",
        len(rows), len(col_keys), file_path,
        ws.title, table_key or "(none)",
    )
    return rows


def detect_sop_table_key(filename):
    """Auto-detect which S&OP table key a filename belongs to.

    Returns one of the SOP_TABLE_KEYS or None.
    """
    fn = filename.lower()

    # DPR / Stock Ledger
    if "stock_ledger" in fn or "stock ledger" in fn or "dpr" in fn or "daily_production" in fn:
        return "sop_dpr"

    # Forecast vs Sales
    if "forecast" in fn:
        return "sop_forecast"

    # Opening Stock / Stock Statement Valuation
    if "stock_statement" in fn or "stock statement" in fn or "opening_stock" in fn or "valuation_report" in fn or "valuation report" in fn:
        return "sop_opening_stock"

    # Green Level / Safety Stock (PPC)
    if "green_level" in fn or "green level" in fn or "safety_stock" in fn:
        return "sop_green_level"

    # Sales Register / Invoice Register
    if "sales_invoice" in fn or "sales invoice" in fn or "sales_register" in fn or "sales register" in fn or "invoice_register" in fn or "invoice register" in fn:
        return "sop_sales_register"

    return None
