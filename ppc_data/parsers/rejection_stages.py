"""Parser for In process-Rejection.xlsx → operation_stage_map (W1.5).

Sheet: "Stage List" (77 ops) + "Process Sheet"
Level: L0
"""

import logging

from openpyxl import load_workbook

from ..field_maps.operation_stage_map import FIELD_MAP, TABLE_KEY
from .base import extract_rows

log = logging.getLogger(__name__)

TABLE_KEY = TABLE_KEY  # noqa: F841


def parse(file_path, sheet_name=None):
    """Parse stage map from rejection file.

    Tries sheets in order: "Stage List", "Process Sheet", first sheet.
    """
    wb = load_workbook(file_path, read_only=True, data_only=True)
    sheets = wb.sheetnames
    wb.close()

    preferred = ["Stage List", "Process Sheet"]
    target = None
    for pref in preferred:
        for name in sheets:
            if pref.lower() in name.lower():
                target = name
                break
        if target:
            break

    rows = extract_rows(file_path, target, FIELD_MAP)
    log.info(
        "Parsed %d stage map rows from %s (sheet: %s)",
        len(rows), file_path, target or "first",
    )
    return rows
