"""Parser for the forecast demand file — Initial Demand sheet (L2).

Source: Forecast file uploaded by PPC.
Supports two formats:
  1. "August forecast 2026.xlsx" style — headers: Part No, Item Code,
     Customer Name, Item Description, Forecast Qty, Revised Qty, R3SS.
     ERP code in "Item Description", demand qty in "R3SS".
  2. Legacy format — headers matching demand_freeze FIELD_MAP directly.

Also handles "Additional W*" and "Reduced Demand" sheets as
demand transactions when present.
"""

import logging

from openpyxl import load_workbook

from ..field_maps.demand_freeze import FIELD_MAP
from .base import clean_cell, detect_header_row

TABLE_KEY = "demand_freeze"

log = logging.getLogger(__name__)

# Sheets that carry the initial demand data
_INITIAL_SHEETS = ["Initial Demand", "Demand", "Forecast", "Sheet1"]

# Sheets that carry additions/reductions
_TX_SHEETS_ADD = ["Additional", "Addition"]
_TX_SHEETS_REDUCE = ["Reduced Demand", "Reduction"]

# Real forecast file column mapping (August forecast style)
_REAL_FORECAST_MAP = {
    "item_description": "Item Description",   # ERP Code → item_code
    "jolly_code":       "Item Code",           # Jolly Code
    "customer":         "Customer Name",
    "r3ss_qty":         "R3SS",                # demand qty
    "forecast_qty":     "Forecast Qty",
    "revised_qty":      "Revised Qty",
}

# Additional sheet column mapping
_REAL_ADDITIONAL_MAP = {
    "item_description": "Item Description",
    "jolly_code":       "Item Code",
    "customer":         "Customer Name",
    "additional_qty":   "Additional Forecast Qty",
    "mpss_qty":         "MPsSS",               # R3SS-used additional qty
}

# Reduced Demand sheet column mapping
_REAL_REDUCED_MAP = {
    "item_description": "Item Description",
    "jolly_code":       "Item Code",
    "customer":         "Customer Name",
    "reduced_qty":      "Reduced Forecast Qty",
}


def _find_sheet(wb, candidates):
    """Find the best sheet matching candidates, prioritizing by candidate order."""
    for cand in candidates:
        for sn in wb.sheetnames:
            if cand.lower() in sn.lower().strip():
                return sn
    return None


def _parse_real_initial(file_path, sheet_name):
    """Parse 'August forecast' style Initial Demand sheet.

    Returns list of dicts with item_code (ERP code) and initial_qty,
    or None if the sheet doesn't match the expected format.
    """
    wb = load_workbook(file_path, data_only=True, read_only=True)
    ws = wb[sheet_name]
    header_map = {v.upper(): k for k, v in _REAL_FORECAST_MAP.items()}
    header_row_idx = None
    col_map = {}

    all_rows = list(ws.iter_rows(values_only=True))
    wb.close()

    for i, row in enumerate(all_rows):
        if i >= 10:
            break
        cells = [str(c).strip() if c is not None else "" for c in row]
        matched = 0
        trial = {}
        for j, cell in enumerate(cells):
            key = header_map.get(cell.upper())
            if key and key not in trial.values():
                trial[j] = key
                matched += 1
        if matched >= 3:
            header_row_idx = i
            col_map = trial
            break

    if header_row_idx is None:
        return None

    rows = []
    for i, row in enumerate(all_rows):
        if i <= header_row_idx:
            continue
        vals = list(row)
        erp_code = None
        qty = 0.0
        for j, key in col_map.items():
            if j >= len(vals):
                continue
            v = vals[j]
            if key == "item_description" and v is not None:
                erp_code = str(v).strip()
            elif key == "r3ss_qty":
                try:
                    qty = float(v) if v is not None else 0.0
                except (ValueError, TypeError):
                    qty = 0.0
        if erp_code and qty > 0:
            rows.append({"item_code": erp_code, "initial_qty": qty})

    return rows if rows else None


def _parse_real_additional(file_path, sheet_name):
    """Parse 'Additional W*' sheet from real forecast file.

    Returns list of dicts with item_code and qty, or empty list.
    """
    wb = load_workbook(file_path, data_only=True, read_only=True)
    ws = wb[sheet_name]
    header_map = {v.upper(): k for k, v in _REAL_ADDITIONAL_MAP.items()}
    header_row_idx = None
    col_map = {}

    all_rows = list(ws.iter_rows(values_only=True))
    wb.close()

    for i, row in enumerate(all_rows):
        if i >= 10:
            break
        cells = [str(c).strip() if c is not None else "" for c in row]
        matched = 0
        trial = {}
        for j, cell in enumerate(cells):
            key = header_map.get(cell.upper())
            if key and key not in trial.values():
                trial[j] = key
                matched += 1
        if matched >= 2:
            header_row_idx = i
            col_map = trial
            break

    if header_row_idx is None:
        return []

    rows = []
    for i, row in enumerate(all_rows):
        if i <= header_row_idx:
            continue
        vals = list(row)
        erp_code = None
        qty = 0.0
        for j, key in col_map.items():
            if j >= len(vals):
                continue
            v = vals[j]
            if key == "item_description" and v is not None:
                erp_code = str(v).strip()
            elif key == "mpss_qty":
                try:
                    qty = float(v) if v is not None else 0.0
                except (ValueError, TypeError):
                    qty = 0.0
        if erp_code and qty > 0:
            rows.append({"item_code": erp_code, "qty": qty})

    return rows


def _parse_real_reduced(file_path, sheet_name):
    """Parse 'Reduced Demand' sheet (Reduced Forecast Qty column)."""
    wb = load_workbook(file_path, data_only=True, read_only=True)
    ws = wb[sheet_name]
    header_map = {v.upper(): k for k, v in _REAL_REDUCED_MAP.items()}

    all_rows = list(ws.iter_rows(values_only=True))
    wb.close()

    header_row_idx = None
    col_map = {}
    for i, row in enumerate(all_rows):
        if i >= 10:
            break
        cells = [str(c).strip() if c is not None else "" for c in row]
        matched = 0
        trial = {}
        for j, cell in enumerate(cells):
            key = header_map.get(cell.upper())
            if key and key not in trial.values():
                trial[j] = key
                matched += 1
        if matched >= 2:
            header_row_idx = i
            col_map = trial
            break

    if header_row_idx is None:
        return []

    rows = []
    for i, row in enumerate(all_rows):
        if i <= header_row_idx:
            continue
        vals = list(row)
        erp_code = None
        qty = 0.0
        for j, key in col_map.items():
            if j >= len(vals):
                continue
            v = vals[j]
            if key == "item_description" and v is not None:
                erp_code = str(v).strip()
            elif key == "reduced_qty":
                try:
                    qty = float(v) if v is not None else 0.0
                except (ValueError, TypeError):
                    qty = 0.0
        if erp_code and qty > 0:
            rows.append({"item_code": erp_code, "qty": qty})

    return rows


def parse(file_path):
    """Parse ALL demand sheets from the forecast file.

    Extracts:
      - Initial Demand sheet → txn_type='INITIAL', qty from 'R3SS' column
      - Additional W* sheets → txn_type='ADDITION', qty from 'MPsSS' column
      - Reduced Demand sheets → txn_type='REDUCTION', qty from that sheet

    Each row has: item_code, qty, txn_type, week (sheet name for additions).
    """
    wb = load_workbook(file_path, data_only=True, read_only=True)

    # Find the initial demand sheet
    sheet_name = _find_sheet(wb, _INITIAL_SHEETS)
    if sheet_name is None:
        sheet_name = wb.sheetnames[0]
        log.warning("No demand sheet found, using first sheet: %s", sheet_name)

    all_rows = []

    # ── 1. Initial Demand ──
    initial = _parse_real_initial(file_path, sheet_name)
    if initial:
        for r in initial:
            all_rows.append({
                "item_code": r["item_code"],
                "initial_qty": r["initial_qty"],
                "qty": r["initial_qty"],
                "txn_type": "INITIAL",
                "week": "",
            })
        log.info("Parsed %d INITIAL rows from %s!%s",
                 len(initial), file_path, sheet_name)
    else:
        # Fall back to legacy FIELD_MAP format
        ws = wb[sheet_name]
        header_to_key = {v.strip(): k for k, v in FIELD_MAP.items()}
        known_headers = list(header_to_key.keys())

        header_row_idx, col_map = detect_header_row(ws, known_headers)

        col_to_key = {}
        for col_idx, header_str in col_map.items():
            key = header_to_key.get(header_str.strip())
            if key:
                col_to_key[col_idx] = key

        for row_idx, raw_row in enumerate(ws.iter_rows(values_only=True)):
            if row_idx <= header_row_idx:
                continue
            record = {"txn_type": "INITIAL", "week": ""}
            any_value = False
            for col_idx, key in col_to_key.items():
                v = clean_cell(raw_row[col_idx]) if col_idx < len(raw_row) else None
                if v is not None:
                    record[key] = v
                    any_value = True
            if any_value and record.get("item_code"):
                all_rows.append(record)

        log.info("Parsed %d INITIAL rows (legacy) from %s!%s",
                 len(all_rows), file_path, sheet_name)

    wb.close()

    # ── 2. Additional & Reduced Demand sheets ──
    txn_rows = parse_transactions(file_path)
    for r in txn_rows:
        tx = r.get("tx_type", "add")
        all_rows.append({
            "item_code": r["item_code"],
            "qty": r["qty"],
            "initial_qty": 0,
            "txn_type": "ADDITION" if tx == "add" else "REDUCTION",
            "week": r.get("week", ""),
        })

    log.info("Parsed %d total demand rows (%d initial + %d transactions) from %s",
             len(all_rows), len(all_rows) - len(txn_rows), len(txn_rows), file_path)
    return all_rows


def parse_transactions(file_path):
    """Parse addition/reduction sheets. Returns list of dicts with an
    extra 'tx_type' key ('add' or 'reduce').
    """
    wb = load_workbook(file_path, data_only=True, read_only=True)
    all_txns = []

    for sn in wb.sheetnames:
        sn_lower = sn.lower().strip()

        tx_type = None
        for cand in _TX_SHEETS_ADD:
            if cand.lower() in sn_lower:
                tx_type = "add"
                break
        if tx_type is None:
            for cand in _TX_SHEETS_REDUCE:
                if cand.lower() in sn_lower:
                    tx_type = "reduce"
                    break
        if tx_type is None:
            continue

        # Try real format (Additional W* with MPsSS column)
        real_rows = _parse_real_additional(file_path, sn)
        if not real_rows and tx_type == "reduce":
            real_rows = _parse_real_reduced(file_path, sn)
        if real_rows:
            for r in real_rows:
                all_txns.append({
                    "tx_type": tx_type,
                    "item_code": r["item_code"],
                    "qty": r["qty"],
                    "week": sn.strip(),
                })
            continue

        # Fall back to legacy format
        from ..field_maps.demand_transaction import FIELD_MAP as TX_MAP
        tx_header_to_key = {v.strip(): k for k, v in TX_MAP.items()}
        known_headers = list(tx_header_to_key.keys())

        ws = wb[sn]
        try:
            header_row_idx, col_map = detect_header_row(ws, known_headers)
        except ValueError:
            log.warning("Skipping sheet %s — no matching headers", sn)
            continue

        col_to_key = {}
        for col_idx, header_str in col_map.items():
            key = tx_header_to_key.get(header_str.strip())
            if key:
                col_to_key[col_idx] = key

        for row_idx, raw_row in enumerate(ws.iter_rows(values_only=True)):
            if row_idx <= header_row_idx:
                continue
            record = {"tx_type": tx_type}
            any_value = False
            for col_idx, key in col_to_key.items():
                v = clean_cell(raw_row[col_idx]) if col_idx < len(raw_row) else None
                if v is not None:
                    record[key] = v
                    any_value = True
            if not record.get("week"):
                record["week"] = sn.strip()
            if any_value and record.get("item_code"):
                all_txns.append(record)

    wb.close()
    log.info("Parsed %d demand transactions from %s", len(all_txns), file_path)
    return all_txns
