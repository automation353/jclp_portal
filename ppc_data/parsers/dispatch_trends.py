"""Parser for Dispatch Trends.xlsx — demand history (L2).

Source: Dispatch Trends.xlsx
Sheets: Demand, Dispatch, Packing, Planning, Data, Sheet4, Master Data.
~5,000 rows of per-part monthly demand/packing/dispatch from 2021.

Strategy: scan all sheets for the one(s) that have recognizable demand
headers. Merge rows from multiple sheets.
"""

import logging

from openpyxl import load_workbook

from ..field_maps.demand_history import FIELD_MAP
from .base import clean_cell, detect_header_row

TABLE_KEY = "demand_history"

log = logging.getLogger(__name__)

# Sheets likely to contain the main data
_PREFERRED = ["Demand", "Data", "Planning"]


def parse(file_path):
    """Parse demand history from Dispatch Trends.xlsx.
    Returns list of dicts keyed by stable field keys.
    """
    wb = load_workbook(file_path, data_only=True, read_only=True)

    header_to_key = {v.strip(): k for k, v in FIELD_MAP.items()}
    known_headers = list(header_to_key.keys())

    all_rows = []

    # Try preferred sheets first, then fall back to all sheets
    sheets_tried = set()
    for name in _PREFERRED:
        for sn in wb.sheetnames:
            if name.lower() in sn.lower() and sn not in sheets_tried:
                sheets_tried.add(sn)
                rows = _extract_sheet(wb[sn], header_to_key, known_headers, sn)
                if rows:
                    all_rows.extend(rows)

    # If nothing found in preferred, try all sheets
    if not all_rows:
        for sn in wb.sheetnames:
            if sn in sheets_tried:
                continue
            rows = _extract_sheet(wb[sn], header_to_key, known_headers, sn)
            if rows:
                all_rows.extend(rows)
                break  # use first sheet that yields rows

    wb.close()
    log.info("Parsed %d demand history rows from %s", len(all_rows), file_path)
    return all_rows


def _extract_sheet(ws, header_to_key, known_headers, sheet_name):
    """Try to extract rows from a single worksheet."""
    try:
        header_row_idx, col_map = detect_header_row(ws, known_headers)
    except ValueError:
        log.debug("Skipping sheet %s — no matching headers", sheet_name)
        return []

    col_to_key = {}
    for col_idx, header_str in col_map.items():
        key = header_to_key.get(header_str.strip())
        if key:
            col_to_key[col_idx] = key

    rows = []
    for row_idx, raw_row in enumerate(ws.iter_rows(values_only=True)):
        if row_idx <= header_row_idx:
            continue
        record = {}
        any_value = False
        for col_idx, key in col_to_key.items():
            v = clean_cell(raw_row[col_idx]) if col_idx < len(raw_row) else None
            if v is not None:
                record[key] = v
                any_value = True
        if any_value and record.get("item_code"):
            rows.append(record)

    return rows
