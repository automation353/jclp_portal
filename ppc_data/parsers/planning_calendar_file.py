"""Parser for MpsSS.xlsm — Mondays of month (L3).

Source: MpsSS.xlsm → "Mondays of month" sheet.
~12 rows. Week-start dates for the planning calendar.
"""

import logging

from openpyxl import load_workbook

from ..field_maps.planning_calendar import FIELD_MAP
from .base import clean_cell, detect_header_row

TABLE_KEY = "planning_calendar"

log = logging.getLogger(__name__)

_CALENDAR_SHEETS = ["Mondays of month", "Mondays", "Calendar"]


def _find_sheet(wb, candidates):
    for sn in wb.sheetnames:
        sn_lower = sn.lower().strip()
        for cand in candidates:
            if cand.lower() in sn_lower:
                return sn
    return None


def parse(file_path):
    """Parse planning calendar from MpsSS.xlsm."""
    wb = load_workbook(file_path, data_only=True, read_only=True)

    sheet_name = _find_sheet(wb, _CALENDAR_SHEETS)
    if sheet_name is None:
        sheet_name = wb.sheetnames[0]
        log.warning("No calendar sheet found, using first: %s", sheet_name)

    ws = wb[sheet_name]

    header_to_key = {v.strip(): k for k, v in FIELD_MAP.items()}
    known_headers = list(header_to_key.keys())

    header_row_idx, col_map = detect_header_row(ws, known_headers)

    col_to_key = {}
    for col_idx, header_str in col_map.items():
        key = header_to_key.get(header_str.strip())
        if key:
            col_to_key[col_idx] = key

    rows = []
    for row_idx, raw_row in enumerate(ws.iter_rows(values_only=True)):
        if row_idx <= header_row_idx:
            continue
        record = {}
        any_value = False
        for col_idx, key in col_to_key.items():
            v = clean_cell(raw_row[col_idx]) if col_idx < len(raw_row) else None
            if v is not None:
                record[key] = v
                any_value = True
        if any_value:
            rows.append(record)

    wb.close()
    log.info("Parsed %d calendar rows from %s!%s", len(rows), file_path, sheet_name)
    return rows
