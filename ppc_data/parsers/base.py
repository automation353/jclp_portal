"""Shared parsing utilities for PPC Excel files.

Mirrors the _clean_cell / _extract_sheet pattern from ppc.api but
generalised for any file type and header-driven column detection.
"""

import logging

from openpyxl import load_workbook

log = logging.getLogger(__name__)


def clean_cell(value):
    """Return a JSON-safe scalar (or None) for a data cell."""
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        try:
            return value.isoformat()
        except Exception:
            return str(value)
    if isinstance(value, (int, float, bool)):
        return value
    s = str(value).strip()
    return s if s else None


def detect_header_row(ws, known_headers, max_rows=10):
    """Scan the first ``max_rows`` rows of a worksheet for the one
    that contains the most matches against ``known_headers``.

    Returns (row_index_0based, {col_index: header_string}).
    Raises ValueError if no row matches even 50% of known headers.
    """
    known_lower = {h.strip().lower() for h in known_headers}
    best_row = None
    best_map = {}
    best_score = 0

    for row_idx, row in enumerate(ws.iter_rows(max_row=max_rows, values_only=True)):
        col_map = {}
        for col_idx, cell in enumerate(row):
            if cell is None:
                continue
            label = str(cell).strip()
            if label.lower() in known_lower:
                col_map[col_idx] = label
        if len(col_map) > best_score:
            best_score = len(col_map)
            best_row = row_idx
            best_map = col_map

    threshold = max(1, len(known_headers) // 2)
    if best_score < threshold:
        raise ValueError(
            f"Could not find a header row matching at least {threshold} of "
            f"{len(known_headers)} expected headers. Best row #{best_row} "
            f"matched only {best_score}: {list(best_map.values())}"
        )
    return best_row, best_map


def extract_rows(file_path, sheet_name, field_map):
    """Open an xlsx, find the header row, and extract all data rows
    re-keyed by stable field keys.

    Args:
        file_path: path to the .xlsx file
        sheet_name: worksheet name (or None for the first sheet)
        field_map: dict of {stable_key: "Excel Header String"}

    Returns:
        list of dicts, each keyed by the stable field keys.
    """
    wb = load_workbook(file_path, data_only=True, read_only=True)

    if sheet_name:
        if sheet_name not in wb.sheetnames:
            available = ", ".join(wb.sheetnames)
            raise ValueError(
                f"Sheet {sheet_name!r} not found. Available: {available}"
            )
        ws = wb[sheet_name]
    else:
        ws = wb.active

    # Build reverse map: Excel header -> stable key
    header_to_key = {v.strip(): k for k, v in field_map.items()}
    known_headers = list(header_to_key.keys())

    header_row_idx, col_map = detect_header_row(ws, known_headers)

    # Build col_index -> stable_key from what we actually found
    col_to_key = {}
    for col_idx, header_str in col_map.items():
        key = header_to_key.get(header_str.strip())
        if key:
            col_to_key[col_idx] = key

    # Extract data rows (everything after the header row)
    rows = []
    for row_idx, raw_row in enumerate(ws.iter_rows(values_only=True)):
        if row_idx <= header_row_idx:
            continue
        record = {}
        any_value = False
        for col_idx, key in col_to_key.items():
            v = clean_cell(raw_row[col_idx]) if col_idx < len(raw_row) else None
            if v is not None:
                record[key] = v
                any_value = True
        if any_value:
            rows.append(record)

    wb.close()
    return rows
