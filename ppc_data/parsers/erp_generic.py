"""Generic parser for manually-uploaded ERP Excel files.

When the TCS iON ERP data is uploaded manually through the PPC upload
screen (instead of via the n8n ERP landing endpoint), this parser handles
the Excel → stable-key conversion using the same ERP field maps.

It tries every sheet in the workbook and picks the one that matches the
most headers — useful when ERP exports have extra cover/summary sheets.
"""

from ..field_maps.erp import ERP_REGISTRY
from .base import clean_cell, detect_header_row

TABLE_KEY = None  # Set dynamically per report_key


def _make_parser(report_key):
    """Build a parser function for a specific ERP report key."""
    entry = ERP_REGISTRY[report_key]
    field_map = entry["field_map"]

    def parse(file_path):
        from openpyxl import load_workbook

        wb = load_workbook(file_path, data_only=True, read_only=True)

        # Build reverse map: Excel header -> stable key
        header_to_key = {v.strip(): k for k, v in field_map.items()}
        known_headers = list(header_to_key.keys())

        # Try each sheet — pick the one with the best header match
        best_ws = None
        best_header_row = None
        best_col_map = None
        best_score = 0

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            try:
                row_idx, col_map = detect_header_row(ws, known_headers)
                if len(col_map) > best_score:
                    best_score = len(col_map)
                    best_ws = ws
                    best_header_row = row_idx
                    best_col_map = col_map
            except ValueError:
                continue

        if best_ws is None:
            wb.close()
            raise ValueError(
                f"No sheet in the workbook matched the expected headers for "
                f"ERP report '{report_key}'. Expected headers like: "
                f"{', '.join(list(known_headers)[:5])}..."
            )

        # Build col_index -> stable_key
        col_to_key = {}
        for col_idx, header_str in best_col_map.items():
            key = header_to_key.get(header_str.strip())
            if key:
                col_to_key[col_idx] = key

        # Extract rows
        rows = []
        for row_idx, raw_row in enumerate(best_ws.iter_rows(values_only=True)):
            if row_idx <= best_header_row:
                continue
            record = {}
            any_value = False
            for col_idx, key in col_to_key.items():
                v = clean_cell(raw_row[col_idx]) if col_idx < len(raw_row) else None
                if v is not None:
                    record[key] = v
                    any_value = True
            if any_value:
                rows.append(record)

        wb.close()
        return rows

    return parse


# Build a parser module-like object for each ERP report key
class _ERPParserProxy:
    """Looks like a parser module: has TABLE_KEY and parse()."""

    def __init__(self, report_key):
        self.TABLE_KEY = report_key
        self.parse = _make_parser(report_key)


# Pre-build all 17 proxies
ERP_PARSERS = {key: _ERPParserProxy(key) for key in ERP_REGISTRY}
